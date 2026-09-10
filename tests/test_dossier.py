from backend.modules.applications.dossier import compact_rows, load_dossier, load_rows, save_rows


def test_compact_merges_generic_courses_under_parent():
    text = compact_rows(
        [
            (
                "Classes préparatoires PCSI-PSI\nContexte : deux ans.",
                "2020-2022 | Nîmes | Lycée Daudet",
            ),
            ("Mathématiques", "Cours, TD, colles, DS"),
            ("Analyse (fonctions, suites, séries)", "Cours, TD, colles, DS"),
            ("Physique", "Cours, TD"),
            (
                "TIPE avalanches\nContexte : modélisation du risque.",
                "2021-2022 | Super Dévoluy",
            ),
            (
                "Développement d'algorithmes de pentes à risque selon un bulletin météo.",
                "Python",
            ),
        ]
    )
    assert "## Classes préparatoires PCSI-PSI" in text
    assert "Matières (ne pas lister une par une sur un CV) : Mathématiques ; Analyse ; Physique" in text
    assert "## Mathématiques" not in text
    assert "Développement d'algorithmes de pentes à risque" in text
    assert "Outils : Python" in text


def test_compact_keeps_project_rows():
    text = compact_rows(
        [
            (
                "Diplôme d'ingénieur en informatique (spécialité IA)",
                "2022-2025 | Lannion | ENSSAT",
            ),
            (
                "JAVA : théorie POO. Projet Quizz solo ou multijoueur.",
                "Java, Swing",
            ),
        ]
    )
    assert "Projet Quizz" in text
    assert "Java, Swing" in text
    assert "Matières" not in text


def test_load_suivi_word_is_compact():
    text = load_dossier()
    assert "outil de configuration industrielle" in text.lower() or "configuration industrielle" in text.lower()
    assert "janvier" in text.lower() or "2026" in text
    assert "Ministère des Transports" in text or "Ministère des transports" in text
    assert "## Mathématiques" not in text
    assert len(text) < 32000


def test_save_rows_roundtrip_updates_word_and_compact_dossier(tmp_path, monkeypatch):
    from docx import Document

    from backend.core.config import settings
    from backend.modules.applications.dossier import _compact_from_docx_cached

    path = tmp_path / "suivi-competences.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Ancienne compétence"
    table.rows[0].cells[1].text = "2020"
    table.rows[1].cells[0].text = "À retirer"
    table.rows[1].cells[1].text = "Java"
    doc.save(path)

    monkeypatch.setattr(settings, "candidate_dossier_path", str(path))
    _compact_from_docx_cached.cache_clear()

    assert load_rows()[0][0] == "Ancienne compétence"

    save_rows(
        [
            ["Outil de configuration industrielle", "C# / .NET | 2026 | IAS"],
            ["", ""],
            ["Mission data mobilité", "Python, Docker"],
        ]
    )

    rows = load_rows()
    assert rows == [
        ["Outil de configuration industrielle", "C# / .NET | 2026 | IAS"],
        ["Mission data mobilité", "Python, Docker"],
    ]
    compact = load_dossier()
    assert "Outil de configuration industrielle" in compact
    assert "Python, Docker" in compact
    assert "Ancienne compétence" not in compact
    assert "À retirer" not in compact


def test_save_rows_can_add_a_third_column(tmp_path, monkeypatch):
    from docx import Document

    from backend.core.config import settings
    from backend.modules.applications.dossier import _compact_from_docx_cached

    path = tmp_path / "suivi-competences.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Projet"
    table.rows[0].cells[1].text = "Python"
    doc.save(path)

    monkeypatch.setattr(settings, "candidate_dossier_path", str(path))
    _compact_from_docx_cached.cache_clear()

    save_rows([["Projet", "Python", "2024 | Lyon"]])
    assert load_rows() == [["Projet", "Python", "2024 | Lyon"]]
    compact = load_dossier()
    assert "Projet" in compact
    assert "2024 | Lyon" in compact
