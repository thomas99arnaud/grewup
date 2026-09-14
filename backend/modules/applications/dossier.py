"""Charge le suivi Word et le compacte pour le LLM (relecture à chaque modification)."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any

from backend.core.config import settings

MODULE_DIR = Path(__file__).resolve().parent
DEFAULT_DOCX = MODULE_DIR / "suivi-competences.docx"
DEFAULT_MD = MODULE_DIR / "dossier.md"

_YEAR = re.compile(r"\b(?:19|20)\d{2}\b")
_PARENS = re.compile(r"\([^)]*\)")
_PROJECT = re.compile(
    r"\b(projet|projets|mandat\s*:|réalisation|tp\s*:|contribution\s*:)",
    re.IGNORECASE,
)
_GENERIC_PART = re.compile(
    r"^(cours|td|colles|ds|tp|python|c|java|lisp|prolog|compte rendu|"
    r"projets|présentations(?: orales)?|pitchs|présentations, pitchs)$",
    re.IGNORECASE,
)


def _norm(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\xa0", " ")).strip()


def _first_line(text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            return stripped
    return ""


def _right_parts(right: str) -> list[str]:
    return [p.strip() for p in re.split(r"[,;/]", right) if p.strip()]


def _is_generic_tools(right: str) -> bool:
    parts = _right_parts(right)
    if not parts:
        return True
    return all(_GENERIC_PART.match(p) for p in parts)


def _is_parent(left: str, right: str) -> bool:
    if _YEAR.search(right):
        return True
    lines = [ln.strip() for ln in left.splitlines() if ln.strip()]
    if len(lines) >= 2 and lines[1].lower().startswith("contexte"):
        return True
    return False


def _course_label(title: str) -> str:
    label = _PARENS.sub(" ", title)
    label = re.split(r"\s[:–—]\s", label, maxsplit=1)[0]
    return _norm(label)[:80]


def _is_course_row(left: str, right: str) -> bool:
    """Matière / titre de cours sans récit de projet — à fusionner."""
    if not left or _is_parent(left, right):
        return False
    if _PROJECT.search(left):
        return False
    if not _is_generic_tools(right):
        return False
    stripped = _norm(_PARENS.sub(" ", left))
    if stripped.endswith("."):
        return False
    return len(stripped) <= 80 and stripped.count(".") == 0


def compact_rows(rows: list[tuple[str, str]]) -> str:
    """Transforme le tableau (gauche, droite) en fiche compacte pour le LLM."""
    lines: list[str] = [
        "# Dossier candidat — Thomas Arnaud",
        "",
        "Source : tableau Word de suivi de compétences. Ne rien inventer au-delà.",
        "",
    ]
    current_courses: list[str] = []

    def flush_courses() -> None:
        nonlocal current_courses
        if not current_courses:
            return
        # Dédupliquer en gardant l'ordre
        seen: set[str] = set()
        unique: list[str] = []
        for title in current_courses:
            key = title.lower()
            if key not in seen:
                seen.add(key)
                unique.append(title)
        lines.append("Matières (ne pas lister une par une sur un CV) : " + " ; ".join(unique))
        lines.append("")
        current_courses = []

    def emit_parent(left: str, right: str) -> None:
        title = _first_line(left)
        body_lines = [ln.strip() for ln in left.splitlines() if ln.strip()]
        rest = body_lines[1:] if body_lines else []
        lines.append(f"## {title}")
        if right:
            lines.append(right)
        if rest:
            lines.append("\n".join(rest))
        lines.append("")

    def emit_fact(left: str, right: str) -> None:
        body = _norm(left)
        if not body:
            return
        lines.append(f"- {body}")
        if right and right.lower() not in body.lower():
            lines.append(f"  Outils : {right}")
        lines.append("")

    for left_raw, right_raw in rows:
        left = left_raw.strip()
        right = right_raw.strip()
        if not left and not right:
            continue
        if _is_parent(left, right):
            flush_courses()
            emit_parent(left, right)
            continue
        if _is_course_row(left, right):
            current_courses.append(_course_label(_first_line(left)))
            continue
        flush_courses()
        emit_fact(left, right)

    flush_courses()
    return "\n".join(lines).strip() + "\n"


def _table_grid(docx_path: Path) -> list[list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "python-docx est requis pour lire le suivi Word. "
            "Installe-le dans le venv : pip install python-docx"
        ) from exc

    table = Document(str(docx_path)).tables[0]
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def _cell_fill(cell) -> str | None:
    from docx.oxml.ns import qn

    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return None
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        return None
    fill = shade.get(qn("w:fill"))
    if not fill or fill.lower() in {"auto", "ffffff"}:
        return None
    return fill.upper()


def _first_run_style(cell) -> tuple[bool, bool, float | None]:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = run.font.size.pt if run.font.size else None
            return bool(run.bold), bool(run.italic), size
    return False, False, None


def _read_cell(cell) -> dict[str, Any]:
    bold, italic, font_size = _first_run_style(cell)
    return {
        "text": cell.text.strip(),
        "fill": _cell_fill(cell),
        "bold": bold,
        "italic": italic,
        "font_size": font_size,
    }


def _table_cells(docx_path: Path) -> list[list[dict[str, Any]]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "python-docx est requis pour lire le suivi Word. "
            "Installe-le dans le venv : pip install python-docx"
        ) from exc

    table = Document(str(docx_path)).tables[0]
    return [[_read_cell(cell) for cell in row.cells] for row in table.rows]


def _pairs_from_grid(grid: list[list[str]]) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for cells in grid:
        if not any(cell.strip() for cell in cells):
            continue
        left = cells[0].strip() if cells else ""
        right = " | ".join(cell.strip() for cell in cells[1:] if cell.strip())
        pairs.append((left, right))
    return pairs


@lru_cache(maxsize=4)
def _compact_from_docx_cached(path_str: str, mtime_ns: int, size: int) -> str:
    return compact_rows(_pairs_from_grid(_table_grid(Path(path_str))))


def _resolve_path() -> Path:
    configured = (settings.candidate_dossier_path or "").strip()
    path = Path(configured) if configured else DEFAULT_DOCX
    if path.exists():
        return path
    if DEFAULT_DOCX.exists() and path != DEFAULT_DOCX:
        return DEFAULT_DOCX
    if DEFAULT_MD.exists():
        return DEFAULT_MD
    raise FileNotFoundError(
        f"Suivi de compétences introuvable : {path}. "
        "Place suivi-competences.docx dans backend/modules/applications/ "
        "ou renseigne CANDIDATE_DOSSIER_PATH."
    )


def load_dossier() -> str:
    """Relit le Word dès qu'il change (mtime) ; compacte les listes de cours."""
    path = _resolve_path()
    if path.suffix.lower() == ".docx":
        stat = path.stat()
        return _compact_from_docx_cached(str(path.resolve()), stat.st_mtime_ns, stat.st_size)
    return path.read_text(encoding="utf-8")


def load_rows() -> list[list[str]]:
    """Toutes les cellules du tableau Word (texte seul)."""
    path = _resolve_path()
    if path.suffix.lower() != ".docx":
        return []
    return _table_grid(path)


def load_cells() -> list[list[dict[str, Any]]]:
    """Cellules du Word avec fond, gras et taille de police."""
    path = _resolve_path()
    if path.suffix.lower() != ".docx":
        return []
    return _table_cells(path)


def dossier_source() -> tuple[str, datetime | None]:
    """Nom du fichier et date de dernière modification."""
    path = _resolve_path()
    mtime = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
    return path.name, mtime


def _docx_write_path() -> Path:
    configured = (settings.candidate_dossier_path or "").strip()
    if configured:
        path = Path(configured)
        if path.suffix.lower() == ".docx":
            return path
    return DEFAULT_DOCX


def _set_cell_fill(cell, fill: str | None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if not fill:
        if shade is not None:
            tc_pr.remove(shade)
        return
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), fill)


def _style_run(run, *, bold: bool, italic: bool, font_size: float | None) -> None:
    from docx.shared import Pt

    run.bold = bold or None
    run.italic = italic or None
    run.font.size = Pt(font_size) if font_size else None


def _write_cell(cell, data: dict[str, Any]) -> None:
    text = data.get("text") or ""
    lines = text.split("\n")
    cell.text = lines[0] if lines else ""
    first = cell.paragraphs[0]
    if not first.runs:
        first.add_run(first.text)
    for run in first.runs:
        _style_run(
            run,
            bold=bool(data.get("bold")),
            italic=bool(data.get("italic")),
            font_size=data.get("font_size"),
        )
    for line in lines[1:]:
        paragraph = cell.add_paragraph(line)
        for run in paragraph.runs:
            _style_run(run, bold=False, italic=False, font_size=None)
    _set_cell_fill(cell, data.get("fill"))


def _empty_cell() -> dict[str, Any]:
    return {"text": "", "fill": None, "bold": False, "italic": False, "font_size": None}


def _as_cell(value: str | dict[str, Any]) -> dict[str, Any]:
    if isinstance(value, str):
        return {**_empty_cell(), "text": value}
    return {
        "text": (value.get("text") or "").strip(),
        "fill": value.get("fill"),
        "bold": bool(value.get("bold")),
        "italic": bool(value.get("italic")),
        "font_size": value.get("font_size"),
    }


def _resize_table(table, needed: int) -> None:
    needed = max(needed, 1)
    while len(table.rows) < needed:
        table.add_row()
    tbl = table._tbl
    while len(table.rows) > needed:
        tbl.remove(table.rows[-1]._tr)


def _normalize_cells(rows: list[list[str | dict[str, Any]]]) -> list[list[dict[str, Any]]]:
    cleaned: list[list[dict[str, Any]]] = []
    for row in rows:
        cells = [_as_cell(cell) for cell in row]
        if any(cell["text"] for cell in cells):
            cleaned.append(cells)
    ncols = max((len(row) for row in cleaned), default=2)
    ncols = max(ncols, 1)
    if not cleaned:
        return [[_empty_cell() for _ in range(ncols)]]
    return [row + [_empty_cell() for _ in range(ncols - len(row))] for row in cleaned]


def _fill_table(table, grid: list[list[dict[str, Any]]]) -> None:
    _resize_table(table, len(grid))
    for index, cells in enumerate(grid):
        row_cells = table.rows[index].cells
        for col, data in enumerate(cells):
            if col < len(row_cells):
                _write_cell(row_cells[col], data)


def _rebuild_first_table(doc, grid: list[list[dict[str, Any]]]) -> None:
    nrows = max(len(grid), 1)
    ncols = max(len(grid[0]) if grid else 0, 1)
    new_table = doc.add_table(rows=nrows, cols=ncols)
    try:
        new_table.style = "Table Grid"
    except ValueError:
        pass
    padded = grid or [[_empty_cell() for _ in range(ncols)]]
    for index, cells in enumerate(padded):
        for col, data in enumerate(cells):
            _write_cell(new_table.cell(index, col), data)
    if len(doc.tables) >= 2:
        old = doc.tables[0]
        old._tbl.addnext(new_table._tbl)
        old._tbl.getparent().remove(old._tbl)


def save_cells(rows: list[list[str | dict[str, Any]]]) -> None:
    """Écrit le tableau Word en conservant fond et police des titres."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "python-docx est requis pour écrire le suivi Word. "
            "Installe-le dans le venv : pip install python-docx"
        ) from exc

    grid = _normalize_cells(rows)
    path = _docx_write_path()
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        doc = Document(str(path))
        if doc.tables:
            table = doc.tables[0]
            same_cols = len(table.columns) == len(grid[0])
            if same_cols:
                _fill_table(table, grid)
            else:
                _rebuild_first_table(doc, grid)
        else:
            _rebuild_first_table(doc, grid)
    else:
        doc = Document()
        _rebuild_first_table(doc, grid)

    doc.save(str(path))
    _compact_from_docx_cached.cache_clear()


def save_rows(rows: list[list[str]]) -> None:
    save_cells(rows)
